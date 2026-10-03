"""In-memory document generators (BytesIO only – nothing is written to disk).

Method statement (DOCX/PDF), BOQ (XLSX/PDF), TDS / PIS PDFs, spec text, cover sheet and the ZIP project pack.
"""
from __future__ import annotations

import io
import zipfile
from datetime import date

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

from utils import data as D
from utils import drawing, logic

PURPLE = colors.HexColor("#52489E")
DEEP = colors.HexColor("#2E2470")
PINK = colors.HexColor("#EF1F51")
GREEN = colors.HexColor("#06AD80")
LAV = colors.HexColor("#F6F4FF")
GREY = colors.HexColor("#6B7280")

WATERMARK_TDS = "ILLUSTRATIVE SAMPLE – NOT AN OFFICIAL TDS"
WATERMARK_DOC = "PROTOTYPE – ILLUSTRATIVE"
FOOTER_TDS = "Prototype for illustration purposes. Refer to official Birla Opus Prime TDS before use."
FOOTER_DOC = "Prototype for illustration purposes only. Products, data and documents are indicative."


def pdf_safe(s) -> str:
    """Helvetica (WinAnsi) can't draw some glyphs – swap them for safe equivalents and escape XML."""
    s = str(s)
    for a, b in {"₹": "Rs ", "≥": ">=", "≤": "<=", "✓": "v", "→": "->", "☐": "[ ]", "—": "-", "✔": "v",
                 "≈": "~", "µ": "u"}.items():
        s = s.replace(a, b)
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _styles():
    ss = getSampleStyleSheet()
    base = ParagraphStyle("base", parent=ss["Normal"], fontName="Helvetica", fontSize=10, leading=14, alignment=TA_LEFT)
    return {
        "base": base,
        "small": ParagraphStyle("small", parent=base, fontSize=8.5, leading=11, textColor=GREY),
        "h1": ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=PURPLE, spaceAfter=6),
        "h2": ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=12.5, leading=16, textColor=DEEP,
                             spaceBefore=10, spaceAfter=4),
        "cell": ParagraphStyle("cell", parent=base, fontSize=8.8, leading=11),
        "cellb": ParagraphStyle("cellb", parent=base, fontName="Helvetica-Bold", fontSize=8.8, leading=11, textColor=colors.white),
    }


def _page_decor(watermark: str, footer: str, header_text: str = "", band: bool = False):
    def draw(canvas, doc):
        w, h = A4
        canvas.saveState()
        if band:
            canvas.setFillColor(PURPLE)
            canvas.rect(0, h - 30 * mm, w, 30 * mm, stroke=0, fill=1)
            canvas.setFillColor(PINK)
            canvas.rect(0, h - 31.5 * mm, w, 1.5 * mm, stroke=0, fill=1)
            canvas.setFillColor(colors.white)
            canvas.setFont("Helvetica-Bold", 9)
            canvas.drawString(18 * mm, h - 9 * mm, "BIRLA OPUS PRIME")
            canvas.setFont("Helvetica-Bold", 17)
            canvas.drawString(18 * mm, h - 19 * mm, pdf_safe(header_text).replace("&amp;", "&")[:70])
        # Watermark
        canvas.setFillColor(colors.Color(0.6, 0.6, 0.6, alpha=0.18))
        canvas.setFont("Helvetica-Bold", 30)
        canvas.translate(w / 2, h / 2)
        canvas.rotate(42)
        canvas.drawCentredString(0, 0, pdf_safe(watermark).replace("&amp;", "&"))
        canvas.restoreState()
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(18 * mm, 10 * mm, pdf_safe(footer).replace("&amp;", "&"))
        canvas.drawRightString(w - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()
    return draw


def _build_pdf(story, watermark=WATERMARK_DOC, footer=FOOTER_DOC, header_text="", band=False) -> bytes:
    buf = io.BytesIO()
    top = 38 * mm if band else 18 * mm
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=top, bottomMargin=18 * mm,
                            title=header_text or "Prime System Builder", author="Prime System Builder (prototype)")
    decor = _page_decor(watermark, footer, header_text, band)
    doc.build(story, onFirstPage=decor, onLaterPages=decor)
    return buf.getvalue()


def _table(rows: list[list], col_widths=None, header=True) -> Table:
    st = _styles()
    data = []
    for r_i, row in enumerate(rows):
        style = st["cellb"] if (header and r_i == 0) else st["cell"]
        data.append([Paragraph(pdf_safe(c), style) for c in row])
    t = Table(data, colWidths=col_widths, repeatRows=1 if header else 0)
    cmds = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D4F5")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1 if header else 0), (-1, -1), [colors.white, LAV])]
    if header:
        cmds.append(("BACKGROUND", (0, 0), (-1, 0), PURPLE))
    t.setStyle(TableStyle(cmds))
    return t


def _context_lines(ctx: dict, system: dict) -> list[tuple[str, str]]:
    site = ctx.get("site", {})
    return [("Persona", ctx.get("persona") or "—"), ("Segment", ctx.get("segment") or "—"),
            ("Surface", ctx.get("surface") or "—"), ("System", system.get("system_name", "—")),
            ("Tier", system.get("tier", "—")), ("Area", f"{system.get('area', 0):,.0f} sq ft"),
            ("Condition / substrate", f"{site.get('condition', '—')} / {site.get('substrate', '—')}"),
            ("City / climate", f"{site.get('city') or '—'} / {site.get('climate', '—')}"),
            ("Warranty (sample)", f"{system.get('warranty', 0):g} years"),
            ("Generated", date.today().strftime("%d %b %Y"))]


# Method statement ---------------------------------------------------------------------
def method_statement_pdf(system: dict, ctx: dict, sections: list[dict]) -> bytes:
    st = _styles()
    story = [Paragraph("Method Statement", st["h1"]),
             Paragraph(pdf_safe(system.get("system_name", "")), st["h2"]),
             _table([[k, v] for k, v in _context_lines(ctx, system)], [50 * mm, 120 * mm], header=False),
             Spacer(1, 6)]
    for n, sec in enumerate(sections, 1):
        block = [Paragraph(f"{n}. {pdf_safe(sec['title'])}", st["h2"])]
        for item in sec["items"]:
            prefix = "[ ] " if sec["title"] == "QA checklist" else "&bull; "
            hold = f' <font color="#F26625"><b>(hold: {pdf_safe(item["hold"])})</b></font>' if item.get("hold") else ""
            block.append(Paragraph(prefix + pdf_safe(item["text"]) + hold, st["base"]))
        story.append(KeepTogether(block[:3]))
        story.extend(block[3:])
    story += [Spacer(1, 10), Paragraph("Sign-off: ____________________ (Applicator) &nbsp;&nbsp; ____________________ (PMC / Engineer)", st["base"])]
    return _build_pdf(story, header_text="Method Statement – " + system.get("system_name", ""), band=True)


def method_statement_docx(system: dict, ctx: dict, sections: list[dict]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    banner = doc.add_paragraph()
    run = banner.add_run("PROTOTYPE – " + FOOTER_DOC)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x8A, 0x64, 0x00)
    h = doc.add_heading("Method Statement", 0)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x52, 0x48, 0x9E)
    p = doc.add_paragraph()
    r = p.add_run(system.get("system_name", ""))
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x2E, 0x24, 0x70)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for k, v in _context_lines(ctx, system):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k, str(v)
    for n, sec in enumerate(sections, 1):
        hd = doc.add_heading(f"{n}. {sec['title']}", level=2)
        for r in hd.runs:
            r.font.color.rgb = RGBColor(0x52, 0x48, 0x9E)
        for item in sec["items"]:
            if sec["title"] == "QA checklist":
                para = doc.add_paragraph("☐ " + item["text"])
            else:
                para = doc.add_paragraph(item["text"], style="List Bullet")
            if item.get("hold"):
                hr = para.add_run(f"  (hold: {item['hold']})")
                hr.bold = True
                hr.font.color.rgb = RGBColor(0xF2, 0x66, 0x25)
    doc.add_paragraph()
    doc.add_paragraph("Sign-off: ____________________ (Applicator)     ____________________ (PMC / Engineer)")
    foot = doc.sections[0].footer.paragraphs[0]
    foot.text = FOOTER_DOC + " Values are SAMPLE."
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# BOQ ---------------------------------------------------------------------------------
def boq_xlsx(df: pd.DataFrame, labour_rate: float, area: float, system: dict, ctx: dict) -> bytes:
    """Excel BOQ with live formulas and brand formatting."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    purple = PatternFill("solid", fgColor="52489E")
    lav = PatternFill("solid", fgColor="F6F4FF")
    amber = PatternFill("solid", fgColor="FFF0C2")
    thin = Side(style="thin", color="D9D4F5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "Bill of Quantities – " + system.get("system_name", "")
    ws["A1"].font = Font(bold=True, size=15, color="52489E")
    ws["A2"] = f"{ctx.get('segment', '')} · {ctx.get('surface', '')} · generated {date.today():%d %b %Y}"
    ws["A2"].font = Font(italic=True, color="6B7280")
    ws["A3"] = "SAMPLE – coverage, pack sizes and rates are illustrative. " + FOOTER_DOC
    ws["A3"].fill = amber
    ws["A3"].font = Font(bold=True, color="8A6400")

    headers = list(logic.BOQ_COLUMNS)
    hr = 5
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.fill, cell.font, cell.border = purple, Font(bold=True, color="FFFFFF"), border
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    # Column letters: A Item, B Code, C Coverage, D Unit, E Coats, F Area, G Wastage, H Qty, I Pack, J Packs, K Rate, L Amount
    r = hr + 1
    for i, row in df.reset_index(drop=True).iterrows():
        vals = [row["Item"], row["Product code"], float(row["Coverage (sq ft/unit)"]), row["Unit"], float(row["Coats"]),
                float(row["Area (sq ft)"]), float(row["Wastage %"])]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v)
        ws.cell(row=r, column=8, value=f"=ROUND(F{r}*E{r}/IF(C{r}>0,C{r},1)*(1+G{r}/100),2)")
        ws.cell(row=r, column=9, value=float(row["Pack size"]))
        ws.cell(row=r, column=10, value=f"=IF(H{r}>0,ROUNDUP(H{r}/IF(I{r}>0,I{r},1),0),0)")
        ws.cell(row=r, column=11, value=float(row["Rate per pack (₹)"]))
        ws.cell(row=r, column=12, value=f"=J{r}*K{r}")
        for c in range(1, 13):
            ws.cell(row=r, column=c).border = border
            if i % 2:
                ws.cell(row=r, column=c).fill = lav
        for c in (11, 12):
            ws.cell(row=r, column=c).number_format = '"₹"#,##,##0'
        r += 1
    last = r - 1
    first = hr + 1
    summary = [("Material total (₹)", f"=SUM(L{first}:L{last})"),
               ("Labour rate (₹/sq ft)", float(labour_rate)),
               ("Area (sq ft)", float(area)),
               ("Labour total (₹)", f"=L{r + 2}*L{r + 3}"),
               ("Grand total (₹)", f"=L{r + 1}+L{r + 4}"),
               ("₹ per sq ft", f"=IF(L{r + 3}>0,L{r + 5}/L{r + 3},0)")]
    for k, (label, val) in enumerate(summary, start=1):
        ws.cell(row=r + k, column=11, value=label).font = Font(bold=True, color="2E2470")
        cell = ws.cell(row=r + k, column=12, value=val)
        cell.font = Font(bold=True, color="EF1F51" if "Grand" in label else "1A1A2E")
        cell.number_format = '#,##0.00' if "per sq ft" in label else '"₹"#,##,##0' if "₹" in label else "#,##0"
        cell.border = border
    widths = [34, 12, 12, 8, 7, 11, 10, 11, 9, 9, 22, 16]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = f"A{hr + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def boq_pdf(df: pd.DataFrame, tot: dict, system: dict, ctx: dict) -> bytes:
    st = _styles()
    rows = [["Item", "Code", "Cov.", "Coats", "Area", "Wst %", "Qty", "Pack", "Packs", "Rate", "Amount"]]
    for _, r in df.iterrows():
        rows.append([r["Item"], r["Product code"], f"{r['Coverage (sq ft/unit)']:g}/{r['Unit']}", f"{r['Coats']:g}",
                     f"{r['Area (sq ft)']:,.0f}", f"{r['Wastage %']:g}", f"{r['Quantity']:,.1f}", f"{r['Pack size']:g}",
                     f"{int(r['Packs needed'])}", logic.inr(r["Rate per pack (₹)"]), logic.inr(r["Amount (₹)"])])
    widths = [42 * mm, 14 * mm, 17 * mm, 10 * mm, 14 * mm, 10 * mm, 14 * mm, 10 * mm, 11 * mm, 16 * mm, 18 * mm]
    story = [Paragraph("Bill of Quantities (sample)", st["h1"]),
             Paragraph(pdf_safe(f"{system.get('system_name', '')} · {ctx.get('segment', '')} · {ctx.get('surface', '')}"), st["h2"]),
             _table(rows, widths), Spacer(1, 10),
             _table([["Summary", "Value"], ["Material", logic.inr(tot["material"])], ["Labour", logic.inr(tot["labour"])],
                     ["Grand total", logic.inr(tot["grand"])], ["Per sq ft", f"Rs {tot['per_sqft']:,.1f}"]],
                    [60 * mm, 50 * mm]),
             Spacer(1, 8), Paragraph("All coverage, pack sizes and rates are SAMPLE / illustrative values.", st["small"])]
    return _build_pdf(story, header_text="BOQ – " + system.get("system_name", ""), band=True)


# TDS / PIS -------------------------------------------------------------------------------
def _val(prod: dict, col: str, suffix: str = "") -> str:
    v = D.text(prod.get(col))
    try:
        v = f"{float(v):g}" if v else v
    except ValueError:
        pass
    return f"{v}{suffix}" if v else "Illustrative value"


def tds_pdf(code: str) -> bytes:
    prod = D.row_dict(D.products(), "code", code)
    st = _styles()
    if not prod:
        story = [Paragraph(f"Product {pdf_safe(code)} not found in products.csv", st["h1"])]
        return _build_pdf(story, WATERMARK_TDS, FOOTER_TDS, code, True)
    name = D.text(prod.get("name"), code)
    surfaces = D.split_pipes(prod.get("surfaces"))
    tech = D.text(prod.get("technology"), "—")
    benefits = [f"Technology: {tech}", f"Elongation: {_val(prod, 'elongation_pct', ' %')}",
                f"Warranty (sample): {_val(prod, 'warranty_yrs', ' years')}", "Part of a complete Prime waterproofing system",
                "Low-odour, water-based" if D.as_bool(prod.get("low_odour", False)) else "Use with adequate ventilation"]
    props = [["Property", "Typical value (illustrative)", "Test method"],
             ["Elongation at break", _val(prod, "elongation_pct", " %"), "ASTM D412"],
             ["Tensile strength", _val(prod, "tensile_mpa", " MPa"), "ASTM D412"],
             ["Crack bridging", _val(prod, "crack_bridging_mm", " mm"), "EN 1062-7"],
             ["Solids content", _val(prod, "solids_pct", " %"), "ASTM D2697"],
             ["VOC", _val(prod, "voc_gpl", " g/L"), "ASTM D3960"],
             ["Recoat time", _val(prod, "recoat_time_hr", " h"), "—"],
             ["Full cure", _val(prod, "full_cure_days", " days"), "—"],
             ["Shelf life", _val(prod, "shelf_life_months", " months"), "—"]]
    cov = D.text(prod.get("coverage_sqft_per_unit"))
    unit = D.text(prod.get("unit"), "unit")
    story = [
        Paragraph(f"Technical Data Sheet – {pdf_safe(code)}", st["h1"]),
        Paragraph(pdf_safe(name), st["h2"]),
        Paragraph("Description", st["h2"]), Paragraph(pdf_safe(D.text(prod.get("description"), "Illustrative value")), st["base"]),
        Paragraph("Uses", st["h2"]), Paragraph(pdf_safe(", ".join(surfaces) or "Illustrative value"), st["base"]),
        Paragraph("Features &amp; benefits", st["h2"]),
        *[Paragraph("&bull; " + pdf_safe(b), st["base"]) for b in benefits],
        Paragraph("Technical properties", st["h2"]), _table(props, [55 * mm, 60 * mm, 55 * mm]),
        Paragraph("Coverage", st["h2"]),
        Paragraph(pdf_safe(f"~{cov} sq ft per {unit} per coat (sample). Recommended coats: {_val(prod, 'coats')}. "
                           "Actual coverage depends on substrate porosity and method." if cov else "Illustrative value"), st["base"]),
        Paragraph("Application method", st["h2"]),
        Paragraph(pdf_safe("Prepare the surface as per the Prime method statement; detail corners and outlets; apply the specified "
                           "coats; cure and test before protection. Refer to the system method statement for full steps."), st["base"]),
        Paragraph("Packaging", st["h2"]), Paragraph(pdf_safe(f"{_val(prod, 'pack_size')} {unit} pack (sample)."), st["base"]),
        Paragraph("Storage", st["h2"]), Paragraph(pdf_safe(f"Store in a cool, dry, shaded place in original sealed packs. Shelf life: "
                                                          f"{_val(prod, 'shelf_life_months', ' months')}."), st["base"]),
        Paragraph("Health &amp; safety", st["h2"]), Paragraph("Wear gloves and goggles. Avoid contact with eyes and skin. Keep out of reach "
                                                              "of children. Refer to the Safety Data Sheet.", st["base"]),
        Paragraph("Standards", st["h2"]), Paragraph(pdf_safe(", ".join(D.split_pipes(prod.get("standards"))) or "—"), st["base"]),
        Spacer(1, 8), Paragraph("All values on this sheet are SAMPLE / illustrative and not an official TDS.", st["small"]),
    ]
    return _build_pdf(story, WATERMARK_TDS, FOOTER_TDS, f"{code} – {name}", band=True)


def pis_pdf(code: str) -> bytes:
    prod = D.row_dict(D.products(), "code", code)
    st = _styles()
    name = D.text(prod.get("name"), code) if prod else code
    rows = [["Field", "Information (illustrative)"],
            ["Product", name], ["Code", code], ["Technology", _val(prod, "technology")],
            ["Use on", ", ".join(D.split_pipes(prod.get("surfaces"))) or "Illustrative value"],
            ["Coverage", f"~{_val(prod, 'coverage_sqft_per_unit')} sq ft/{D.text(prod.get('unit'), 'unit')}/coat"],
            ["Coats", _val(prod, "coats")], ["Pack", f"{_val(prod, 'pack_size')} {D.text(prod.get('unit'), '')}"],
            ["Warranty (sample)", _val(prod, "warranty_yrs", " years")], ["Standards", ", ".join(D.split_pipes(prod.get("standards"))) or "—"]]
    story = [Paragraph(f"Product Information Sheet – {pdf_safe(code)}", st["h1"]),
             Paragraph(pdf_safe(D.text(prod.get("description"), "Illustrative value")), st["base"]), Spacer(1, 8),
             _table(rows, [45 * mm, 125 * mm])]
    return _build_pdf(story, WATERMARK_TDS, FOOTER_TDS, f"PIS – {code}", band=True)


def pdf_preview_png(pdf_bytes: bytes, zoom: float = 1.3) -> bytes | None:
    try:
        import pymupdf as fitz

        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        return doc[0].get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False).tobytes("png")
    except Exception:  # noqa: BLE001
        return None


# Cover sheet + ZIP pack ----------------------------------------------------------------------
def cover_sheet_pdf(system: dict, ctx: dict) -> bytes:
    st = _styles()
    layers = [["#", "Layer", "Product / item", "Coats"]] + [
        [str(i), l["layer_name"], f"{l['code']} – {l['name']}" if l["is_product"] else f"{l['name']} (generic)", str(l["coats"])]
        for i, l in enumerate(system.get("layers", []), 1)]
    story = [Paragraph("Project Pack – Cover Sheet", st["h1"]),
             Paragraph(pdf_safe(system.get("system_name", "")), st["h2"]),
             _table([["Selection", "Value"]] + [[k, v] for k, v in _context_lines(ctx, system)], [55 * mm, 115 * mm]),
             Paragraph("System layers (bottom to top)", st["h2"]), _table(layers, [10 * mm, 50 * mm, 95 * mm, 15 * mm]),
             Paragraph("Contents of this pack", st["h2"]),
             *[Paragraph("&bull; " + x, st["base"]) for x in
               ["01_Cover_Sheet.pdf", "02_Method_Statement.pdf", "03_BOQ.xlsx", "04_TDS_*.pdf (illustrative)",
                "05_Section_Detail.dxf", "06_Specification.txt"]],
             Spacer(1, 10), Paragraph(pdf_safe(FOOTER_DOC + " All values are SAMPLE."), st["small"])]
    return _build_pdf(story, header_text="Prime System Builder – Project Pack", band=True)


def project_pack_zip(system: dict, ctx: dict, sections: list[dict], boq_df: pd.DataFrame, labour_rate: float,
                     spec: str, substrate: str) -> bytes:
    buf = io.BytesIO()
    area = system.get("area", 0) or 0
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("01_Cover_Sheet.pdf", cover_sheet_pdf(system, ctx))
        z.writestr("02_Method_Statement.pdf", method_statement_pdf(system, ctx, sections))
        z.writestr("03_BOQ.xlsx", boq_xlsx(boq_df, labour_rate, area, system, ctx))
        for code in sorted({l["code"] for l in system.get("layers", []) if l["is_product"]}):
            z.writestr(f"04_TDS_{code.replace('+', 'plus')}.pdf", tds_pdf(code))
        z.writestr("05_Section_Detail.dxf", drawing.section_dxf(system.get("layers", []), substrate, system.get("system_name", "")))
        z.writestr("06_Specification.txt", spec)
        z.writestr("README.txt", "Prime System Builder – project pack (prototype).\n" + FOOTER_DOC + "\nAll values are SAMPLE.\n")
    return buf.getvalue()


def all_tds_zip(codes: list[str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for code in codes:
            safe = code.replace("+", "plus")
            z.writestr(f"TDS_{safe}.pdf", tds_pdf(code))
            z.writestr(f"PIS_{safe}.pdf", pis_pdf(code))
    return buf.getvalue()
